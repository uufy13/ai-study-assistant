"""PPTX 文件解析器（python-pptx）。"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn

from ._common import check_ooxml_zip, iso_datetime, make_markdown_table, normalize_text
from .document import Document
from .errors import FileReadError, InvalidDocumentError


class PptxParser:
    """解析 .pptx 文件，输出带幻灯片分隔的 Markdown 风格文本 Document。

    每张幻灯片提取标题、正文文本、表格与演讲者备注；图片/图表/动画跳过。
    演讲者备注按设计实现，但需真实 fixture 测试验证后才能标记为稳定支持。
    """

    def parse(self, path: Path) -> Document:
        if not path.exists():
            raise FileReadError(f"File not found: {path}")

        check_ooxml_zip(path, "PPTX")

        try:
            prs = Presentation(str(path))
        except Exception as exc:  # python-pptx 对损坏文件的异常种类不统一
            raise InvalidDocumentError(f"Failed to parse PPTX: {path}") from exc

        return self._build(path, prs)

    def _build(self, path: Path, prs) -> Document:
        slide_count = len(prs.slides)
        blocks: list[str] = []
        table_count = 0
        has_notes = False

        for index, slide in enumerate(prs.slides):
            blocks.append(f"--- Slide {index + 1} ---")
            for shape in slide.shapes:
                if getattr(shape, "has_table", False):
                    rendered = self._render_table(shape.table)
                    if rendered:
                        blocks.append(rendered)
                        table_count += 1
                elif getattr(shape, "has_text_frame", False):
                    rendered = self._render_text_shape(shape)
                    if rendered:
                        blocks.append(rendered)
            notes = _extract_notes(slide)
            if notes:
                blocks.append("--- Speaker Notes ---")
                blocks.append(notes)
                has_notes = True

        content = normalize_text("\n\n".join(blocks))
        metadata = _build_metadata(prs, slide_count, has_notes, table_count)
        return Document(source=path, format="pptx", content=content, metadata=metadata)

    def _render_text_shape(self, shape) -> str:
        text = shape.text_frame.text.strip()
        if not text:
            return ""
        if _is_title(shape):
            return "# " + text
        return _render_text_frame(shape.text_frame)

    @staticmethod
    def _render_table(table) -> str:
        rows: list[list[str]] = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        return make_markdown_table(rows)


def _render_text_frame(text_frame) -> str:
    """把文本框的段落展平为 bullet 列表，层级用缩进表示，编号列表用 ``1. ``。"""
    lines: list[str] = []
    for para in text_frame.paragraphs:
        text = "".join(run.text for run in para.runs).strip()
        if not text:
            continue
        indent = "  " * (para.level or 0)
        marker = "1. " if _is_ordered(para) else "- "
        lines.append(indent + marker + text)
    return "\n".join(lines)


def _is_ordered(para) -> bool:
    """段落是否为有序（编号）列表（``a:buAutoNum``/``a:buNum``）。"""
    p_pr = para._p.find(qn("a:pPr"))
    if p_pr is None:
        return False
    return p_pr.find(qn("a:buAutoNum")) is not None or p_pr.find(qn("a:buNum")) is not None


def _is_title(shape) -> bool:
    if not getattr(shape, "is_placeholder", False):
        return False
    try:
        return shape.placeholder_format.type == PP_PLACEHOLDER.TITLE
    except Exception:  # noqa: BLE001 - 占位符类型读取失败按非标题处理
        return False


def _extract_notes(slide) -> str:
    try:
        if slide.has_notes_slide:
            return slide.notes_slide.notes_text_frame.text.strip()
    except Exception:  # noqa: BLE001 - 备注读取失败不致命
        return ""
    return ""


def _build_metadata(prs, slide_count: int, has_notes: bool, table_count: int) -> dict:
    metadata: dict = {
        "parser_name": "pptx",
        "parser_version": "0.2.0",
        "slide_count": slide_count,
        "has_notes": has_notes,
        "table_count": table_count,
    }
    try:
        cp = prs.core_properties
        fields = {
            "title": cp.title,
            "author": cp.author,
            "created_at": iso_datetime(cp.created),
            "modified_at": iso_datetime(cp.modified),
        }
    except Exception:  # noqa: BLE001 - 核心属性缺失时不致命
        fields = {}
    for key, value in fields.items():
        if value:
            metadata[key] = value
    return metadata
