"""PPTX Parser 单元测试（虚构、脱敏、最小化 fixture）。"""

import struct
import zipfile
import zlib
from datetime import UTC, datetime
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from ai_study_assistant.parsers.errors import (
    FileReadError,
    InvalidDocumentError,
    PasswordProtectedError,
)
from ai_study_assistant.parsers.pptx_parser import PptxParser


def _make_png(width: int = 1, height: int = 1) -> bytes:
    """用标准库构造最小合法 PNG（避免引入 Pillow）。"""

    def chunk(typ: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + typ
            + data
            + struct.pack(">I", zlib.crc32(typ + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 6, 0, 0, 0)  # 8-bit RGBA
    raw = b"\x00" + b"\x00\x00\x00\x00" * width
    idat = zlib.compress(raw * height)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", idat)
        + chunk(b"IEND", b"")
    )


def _build_sample(prs: Presentation) -> None:
    # Slide 1：标题 + 要点
    slide1 = prs.slides.add_slide(prs.slide_layouts[1])
    slide1.shapes.title.text = "演示标题"
    body = slide1.placeholders[1].text_frame
    body.text = "要点一"
    body.add_paragraph().text = "要点二"

    # Slide 2：标题 + 演讲者备注
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = "第二页"
    slide2.notes_slide.notes_text_frame.text = "这是备注内容"

    # Slide 3：表格
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide3.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
    shape.table.cell(0, 0).text = "A"
    shape.table.cell(0, 1).text = "B"
    shape.table.cell(1, 0).text = "1"
    shape.table.cell(1, 1).text = "2"


def _save(path: Path, builder) -> Path:
    prs = Presentation()
    builder(prs)
    prs.save(str(path))
    return path


def test_parse_pptx_basic(tmp_path: Path) -> None:
    path = _save(tmp_path / "sample.pptx", _build_sample)
    doc = PptxParser().parse(path)

    assert doc.format == "pptx"
    assert "--- Slide 1 ---" in doc.content
    assert "# 演示标题" in doc.content
    assert "- 要点一" in doc.content
    assert "- 要点二" in doc.content
    assert "--- Slide 2 ---" in doc.content
    assert "# 第二页" in doc.content
    assert "--- Speaker Notes ---" in doc.content
    assert "这是备注内容" in doc.content
    assert "--- Slide 3 ---" in doc.content
    assert "| A | B |" in doc.content
    assert "| 1 | 2 |" in doc.content
    # 仅 Slide 2 有备注，故只出现一次分隔符
    assert doc.content.count("--- Speaker Notes ---") == 1
    # 表格单元格不重复输出：不独立成行
    lines = [line.strip() for line in doc.content.split("\n")]
    assert "A" not in lines
    assert "1" not in lines


def test_parse_pptx_slide_order(tmp_path: Path) -> None:
    def builder(prs: Presentation) -> None:
        for i in range(1, 4):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = f"第{i}页标题"

    path = _save(tmp_path / "order.pptx", builder)
    doc = PptxParser().parse(path)

    content = doc.content
    assert content.index("--- Slide 1 ---") < content.index("第1页标题")
    assert content.index("第1页标题") < content.index("--- Slide 2 ---")
    assert content.index("--- Slide 2 ---") < content.index("第2页标题")
    assert content.index("第2页标题") < content.index("--- Slide 3 ---")
    assert content.index("--- Slide 3 ---") < content.index("第3页标题")
    assert doc.metadata["slide_count"] == 3


def test_parse_pptx_ordered_list(tmp_path: Path) -> None:
    def builder(prs: Presentation) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "列表标题"
        body = slide.placeholders[1].text_frame
        body.text = "无序要点"
        p2 = body.add_paragraph()
        p2.text = "有序一"
        p_pr = p2._p.get_or_add_pPr()
        p_pr.append(p_pr.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"}))

    path = _save(tmp_path / "ordered.pptx", builder)
    doc = PptxParser().parse(path)

    assert "- 无序要点" in doc.content
    assert "1. 有序一" in doc.content


def test_parse_pptx_notes_binding(tmp_path: Path) -> None:
    def builder(prs: Presentation) -> None:
        s1 = prs.slides.add_slide(prs.slide_layouts[5])
        s1.shapes.title.text = "第一页"
        s1.notes_slide.notes_text_frame.text = "第一页备注"
        s2 = prs.slides.add_slide(prs.slide_layouts[5])
        s2.shapes.title.text = "第二页"
        s2.notes_slide.notes_text_frame.text = "第二页备注"

    path = _save(tmp_path / "notes.pptx", builder)
    doc = PptxParser().parse(path)

    content = doc.content
    # 第一页备注绑定在第一页与第二页之间，未与第二页备注串联
    assert content.index("--- Slide 1 ---") < content.index("第一页备注")
    assert content.index("第一页备注") < content.index("--- Slide 2 ---")
    assert content.index("--- Slide 2 ---") < content.index("第二页备注")
    # 每页备注各有一个独立分隔符
    assert content.count("--- Speaker Notes ---") == 2
    assert content.count("第一页备注") == 1
    assert content.count("第二页备注") == 1
    assert doc.metadata["has_notes"] is True


def test_parse_pptx_no_notes(tmp_path: Path) -> None:
    def builder(prs: Presentation) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "无备注标题"

    path = _save(tmp_path / "no_notes.pptx", builder)
    doc = PptxParser().parse(path)

    assert "--- Speaker Notes ---" not in doc.content
    assert doc.metadata["has_notes"] is False


def test_parse_pptx_metadata(tmp_path: Path) -> None:
    path = _save(tmp_path / "sample.pptx", _build_sample)
    doc = PptxParser().parse(path)

    assert doc.metadata["parser_name"] == "pptx"
    assert doc.metadata["parser_version"] == "0.2.0"
    assert doc.metadata["slide_count"] == 3
    assert doc.metadata["has_notes"] is True
    assert doc.metadata["table_count"] == 1


def test_parse_pptx_metadata_full(tmp_path: Path) -> None:
    def builder(prs: Presentation) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "内容"
        prs.core_properties.title = "虚构标题"
        prs.core_properties.author = "虚构作者"
        prs.core_properties.created = datetime(2024, 1, 2, 3, 4, 5, tzinfo=UTC)
        prs.core_properties.modified = datetime(2024, 6, 7, 8, 9, 10, tzinfo=UTC)

    path = _save(tmp_path / "meta.pptx", builder)
    doc = PptxParser().parse(path)

    assert doc.metadata["parser_name"] == "pptx"
    assert doc.metadata["parser_version"] == "0.2.0"
    assert doc.metadata["slide_count"] == 1
    assert doc.metadata["has_notes"] is False
    assert doc.metadata["table_count"] == 0
    assert doc.metadata["title"] == "虚构标题"
    assert doc.metadata["author"] == "虚构作者"
    assert doc.metadata["created_at"] == "2024-01-02T03:04:05"
    assert doc.metadata["modified_at"] == "2024-06-07T08:09:10"


def test_parse_pptx_empty(tmp_path: Path) -> None:
    path = _save(tmp_path / "empty.pptx", lambda prs: None)
    doc = PptxParser().parse(path)

    assert doc.format == "pptx"
    assert doc.content == ""
    assert doc.metadata["slide_count"] == 0


def test_parse_pptx_unicode(tmp_path: Path) -> None:
    def builder(prs: Presentation) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "中文标题 😀"

    path = _save(tmp_path / "unicode.pptx", builder)
    doc = PptxParser().parse(path)

    assert "# 中文标题" in doc.content
    assert "😀" in doc.content


def test_parse_pptx_image_skipped(tmp_path: Path) -> None:
    def builder(prs: Presentation) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = "图片说明文字"
        png = tmp_path / "pixel.png"
        png.write_bytes(_make_png())
        slide.shapes.add_picture(str(png), Inches(1), Inches(2), Inches(2), Inches(2))

    path = _save(tmp_path / "image.pptx", builder)
    doc = PptxParser().parse(path)

    assert "图片说明文字" in doc.content
    assert "base64" not in doc.content.lower()
    assert "\x89PNG" not in doc.content


def test_parse_pptx_chart_skipped(tmp_path: Path) -> None:
    def builder(prs: Presentation) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = "图表说明文字"
        chart_data = CategoryChartData()
        chart_data.categories = ["类别A", "类别B"]
        chart_data.add_series("系列1", (1, 2))
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(1),
            Inches(2),
            Inches(4),
            Inches(3),
            chart_data,
        )

    path = _save(tmp_path / "chart.pptx", builder)
    doc = PptxParser().parse(path)

    assert "图表说明文字" in doc.content
    # 图表不解析嵌入 Excel 数据：类别/系列名称不进入内容
    assert "类别A" not in doc.content
    assert "系列1" not in doc.content


def test_parse_pptx_corrupted(tmp_path: Path) -> None:
    path = tmp_path / "corrupted.pptx"
    path.write_text("not a real pptx", encoding="utf-8")

    with pytest.raises(InvalidDocumentError):
        PptxParser().parse(path)


def test_parse_pptx_encrypted(tmp_path: Path) -> None:
    path = tmp_path / "encrypted.pptx"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("EncryptionInfo", "fake-encryption")

    with pytest.raises(PasswordProtectedError):
        PptxParser().parse(path)


def test_parse_pptx_missing_file(tmp_path: Path) -> None:
    with pytest.raises(FileReadError):
        PptxParser().parse(tmp_path / "missing.pptx")
