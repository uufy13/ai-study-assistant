"""DOCX Parser 单元测试（虚构、脱敏、最小化 fixture）。"""

import struct
import zipfile
import zlib
from datetime import UTC, datetime
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ai_study_assistant.parsers.docx_parser import DocxParser
from ai_study_assistant.parsers.errors import (
    FileReadError,
    InvalidDocumentError,
    PasswordProtectedError,
)


def _make_png(width: int = 1, height: int = 1) -> bytes:
    """用标准库构造最小合法 PNG（避免引入 Pillow / 依赖外部 fixture 文件）。"""

    def chunk(typ: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + typ
            + data
            + struct.pack(">I", zlib.crc32(typ + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 6, 0, 0, 0)  # 8-bit RGBA
    raw = b"\x00" + b"\x00\x00\x00\x00" * width  # 每行 filter=0 + RGBA 像素
    idat = zlib.compress(raw * height)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", idat)
        + chunk(b"IEND", b"")
    )


def _build_sample(doc: DocxDocument) -> None:
    doc.add_heading("标题一", level=1)
    doc.add_heading("小节", level=2)
    doc.add_paragraph("这是正文段落。")
    doc.add_paragraph("项目一", style="List Bullet")
    doc.add_paragraph("项目二", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    table.rows[1].cells[0].text = "1"
    table.rows[1].cells[1].text = "2"


def _save(path: Path, builder) -> Path:
    doc = DocxDocument()
    builder(doc)
    doc.save(str(path))
    return path


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    """向段落追加外部超链接 run（python-docx 无内置 API，手工构造 w:hyperlink）。"""
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def test_parse_docx_basic(tmp_path: Path) -> None:
    path = _save(tmp_path / "sample.docx", _build_sample)
    doc = DocxParser().parse(path)

    assert doc.format == "docx"
    assert "# 标题一" in doc.content
    assert "## 小节" in doc.content
    assert "这是正文段落。" in doc.content
    assert "- 项目一" in doc.content
    assert "- 项目二" in doc.content
    assert "| A | B |" in doc.content
    assert "| 1 | 2 |" in doc.content


def test_parse_docx_heading_levels(tmp_path: Path) -> None:
    def builder(doc: DocxDocument) -> None:
        doc.add_heading("一级", level=1)
        doc.add_heading("二级", level=2)
        doc.add_heading("三级", level=3)

    path = _save(tmp_path / "headings.docx", builder)
    doc = DocxParser().parse(path)

    assert "# 一级" in doc.content
    assert "## 二级" in doc.content
    assert "### 三级" in doc.content


def test_parse_docx_ordered_list(tmp_path: Path) -> None:
    def builder(doc: DocxDocument) -> None:
        doc.add_paragraph("第一项", style="List Number")
        doc.add_paragraph("第二项", style="List Number")

    path = _save(tmp_path / "ordered.docx", builder)
    doc = DocxParser().parse(path)

    assert "1. 第一项" in doc.content
    assert "1. 第二项" in doc.content


def test_parse_docx_nested_lists(tmp_path: Path) -> None:
    def builder(doc: DocxDocument) -> None:
        doc.add_paragraph("外层无序", style="List Bullet")
        doc.add_paragraph("内层无序", style="List Bullet 2")
        doc.add_paragraph("外层有序", style="List Number")
        doc.add_paragraph("内层有序", style="List Number 2")

    path = _save(tmp_path / "nested.docx", builder)
    doc = DocxParser().parse(path)

    # 嵌套列表展平：各层级用统一标记，不保留缩进。
    assert "- 外层无序" in doc.content
    assert "- 内层无序" in doc.content
    assert "1. 外层有序" in doc.content
    assert "1. 内层有序" in doc.content
    assert "  - " not in doc.content


def test_parse_docx_hyperlink(tmp_path: Path) -> None:
    def builder(doc: DocxDocument) -> None:
        p = doc.add_paragraph("访问 ")
        _add_hyperlink(p, "https://example.com/docs", "官方文档")
        p.add_run(" 了解详情")

    path = _save(tmp_path / "hyperlink.docx", builder)
    doc = DocxParser().parse(path)

    assert "[官方文档](https://example.com/docs)" in doc.content
    assert "访问 [官方文档](https://example.com/docs) 了解详情" in doc.content


def test_parse_docx_document_order(tmp_path: Path) -> None:
    def builder(doc: DocxDocument) -> None:
        doc.add_paragraph("第一段")
        t1 = doc.add_table(rows=1, cols=1)
        t1.rows[0].cells[0].text = "表一"
        doc.add_paragraph("第二段")
        t2 = doc.add_table(rows=1, cols=1)
        t2.rows[0].cells[0].text = "表二"
        doc.add_paragraph("第三段")

    path = _save(tmp_path / "order.docx", builder)
    doc = DocxParser().parse(path)

    content = doc.content
    assert content.index("第一段") < content.index("表一")
    assert content.index("表一") < content.index("第二段")
    assert content.index("第二段") < content.index("表二")
    assert content.index("表二") < content.index("第三段")


def test_parse_docx_markdown_table(tmp_path: Path) -> None:
    def builder(doc: DocxDocument) -> None:
        table = doc.add_table(rows=3, cols=3)
        table.rows[0].cells[0].text = "姓名"
        table.rows[0].cells[1].text = "年龄"
        table.rows[0].cells[2].text = "城市"
        table.rows[1].cells[0].text = "张三"
        table.rows[1].cells[1].text = "20"
        table.rows[1].cells[2].text = "北京"
        table.rows[2].cells[0].text = "李四"
        table.rows[2].cells[1].text = "30"
        table.rows[2].cells[2].text = "上海"

    path = _save(tmp_path / "table.docx", builder)
    doc = DocxParser().parse(path)

    assert "| 姓名 | 年龄 | 城市 |" in doc.content
    assert "| --- | --- | --- |" in doc.content
    assert "| 张三 | 20 | 北京 |" in doc.content
    assert "| 李四 | 30 | 上海 |" in doc.content


def test_parse_docx_metadata(tmp_path: Path) -> None:
    def builder(doc: DocxDocument) -> None:
        doc.core_properties.author = "虚构作者"
        doc.core_properties.title = "虚构标题"
        doc.core_properties.created = datetime(2024, 1, 2, 3, 4, 5, tzinfo=UTC)
        doc.core_properties.modified = datetime(2024, 6, 7, 8, 9, 10, tzinfo=UTC)
        doc.add_paragraph("内容")
        doc.add_table(rows=1, cols=1)
        doc.add_table(rows=1, cols=1)

    path = _save(tmp_path / "meta.docx", builder)
    doc = DocxParser().parse(path)

    assert doc.metadata["parser_name"] == "docx"
    assert doc.metadata["parser_version"] == "0.2.0"
    assert doc.metadata["table_count"] == 2
    assert doc.metadata["author"] == "虚构作者"
    assert doc.metadata["title"] == "虚构标题"
    assert doc.metadata["created_at"] == "2024-01-02T03:04:05+00:00"
    assert doc.metadata["modified_at"] == "2024-06-07T08:09:10+00:00"


def test_parse_docx_empty(tmp_path: Path) -> None:
    path = _save(tmp_path / "empty.docx", lambda doc: None)
    doc = DocxParser().parse(path)

    assert doc.format == "docx"
    assert doc.content == ""


def test_parse_docx_unicode(tmp_path: Path) -> None:
    def builder(doc: DocxDocument) -> None:
        doc.add_paragraph("中文内容 😀 全角：ＡＢＣ")

    path = _save(tmp_path / "unicode.docx", builder)
    doc = DocxParser().parse(path)

    assert "中文内容" in doc.content
    assert "😀" in doc.content
    assert "ＡＢＣ" in doc.content


def test_parse_docx_image_skipped(tmp_path: Path) -> None:
    png = tmp_path / "pixel.png"
    png.write_bytes(_make_png())

    def builder(doc: DocxDocument) -> None:
        doc.add_paragraph("图片前")
        doc.add_picture(str(png))
        doc.add_paragraph("图片后")

    path = _save(tmp_path / "image.docx", builder)
    doc = DocxParser().parse(path)

    assert "图片前" in doc.content
    assert "图片后" in doc.content


def test_parse_docx_header_footer_skipped(tmp_path: Path) -> None:
    def builder(doc: DocxDocument) -> None:
        doc.add_paragraph("正文段落")
        doc.sections[0].header.paragraphs[0].text = "页眉不应出现"
        doc.sections[0].footer.paragraphs[0].text = "页脚不应出现"

    path = _save(tmp_path / "header_footer.docx", builder)
    doc = DocxParser().parse(path)

    assert "正文段落" in doc.content
    assert "页眉不应出现" not in doc.content
    assert "页脚不应出现" not in doc.content


def test_parse_docx_corrupted(tmp_path: Path) -> None:
    path = tmp_path / "corrupted.docx"
    path.write_text("not a real docx", encoding="utf-8")

    with pytest.raises(InvalidDocumentError):
        DocxParser().parse(path)


def test_parse_docx_encrypted(tmp_path: Path) -> None:
    path = tmp_path / "encrypted.docx"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("EncryptionInfo", "fake-encryption")

    with pytest.raises(PasswordProtectedError):
        DocxParser().parse(path)


def test_parse_docx_missing_file(tmp_path: Path) -> None:
    with pytest.raises(FileReadError):
        DocxParser().parse(tmp_path / "missing.docx")
