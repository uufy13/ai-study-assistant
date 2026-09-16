"""端到端集成测试：真实 DOCX/PDF/PPTX → parse_file → StudyPipeline → StudyNote → Storage。

所有 LLM 调用使用 FakeLLMClient，禁止真实 API/网络。
"""

from pathlib import Path
from typing import Any

import pymupdf
import pytest
from docx import Document as _DocxDocument
from pptx import Presentation
from pptx.util import Inches

from ai_study_assistant.cli import main
from ai_study_assistant.parsers import parse_file
from ai_study_assistant.pipeline import StudyPipeline
from ai_study_assistant.storage import save


class FakeLLMClient:
    """记录调用并返回固定文本的纯 Python 客户端。"""

    def __init__(self, response: str = "## 生成章节") -> None:
        self.response = response
        self.calls: list[Any] = []

    def generate(self, messages: list[Any], **overrides: Any) -> str:
        self.calls.append(messages)
        return self.response


class _FakeConfig:
    """供 CLI 测试注入：resolve 返回 None，create_client 已被 monkeypatch。"""

    def resolve(self, name: str | None = None) -> None:
        return None


def _patch_cli(monkeypatch: pytest.MonkeyPatch, response: str = "## 生成章节") -> None:
    """替换 load_config/create_client，使 main() 走真实解析但用 fake LLM。"""
    monkeypatch.setattr(
        "ai_study_assistant.cli.load_config", lambda *a, **k: _FakeConfig()
    )
    monkeypatch.setattr(
        "ai_study_assistant.cli.create_client", lambda resolved: FakeLLMClient(response)
    )


# ---------------------------------------------------------------------------
# Fixture builders（真实、脱敏、最小化）
# ---------------------------------------------------------------------------


def _make_docx(path: Path) -> None:
    doc = _DocxDocument()
    doc.add_heading("测试标题", level=1)
    doc.add_paragraph("普通段落内容")
    doc.add_paragraph("列表项一", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    doc.save(str(path))


def _make_pdf(path: Path) -> None:
    doc = pymupdf.open()
    page1 = doc.new_page()
    page1.insert_textbox(pymupdf.Rect(72, 72, 500, 200), "第一页内容", fontname="china-s")
    page2 = doc.new_page()
    page2.insert_textbox(pymupdf.Rect(72, 72, 500, 200), "第二页内容", fontname="china-s")
    doc.save(str(path))
    doc.close()


def _make_pptx(path: Path) -> None:
    prs = Presentation()
    slide1 = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    slide1.shapes.title.text = "第一页标题"
    slide1.placeholders[1].text_frame.text = "要点一"
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    table = slide2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
    table.table.cell(0, 0).text = "A"
    table.table.cell(0, 1).text = "B"
    table.table.cell(1, 0).text = "1"
    table.table.cell(1, 1).text = "2"
    prs.save(str(path))


# ---------------------------------------------------------------------------
# Pipeline 集成：parse_file → Document → StudyPipeline.run → StudyNote → save
# ---------------------------------------------------------------------------


def test_e2e_docx(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    _make_docx(path)

    doc = parse_file(path)
    assert doc.format == "docx"
    # 标题 / 段落 / 列表 / 表格进入统一 Markdown 内容
    assert "# 测试标题" in doc.content
    assert "普通段落内容" in doc.content
    assert "- 列表项一" in doc.content
    assert "| A | B |" in doc.content
    # 无重复表格：单元格文本不独立成行
    lines = [line.strip() for line in doc.content.split("\n")]
    assert "A" not in lines

    client = FakeLLMClient("## 生成章节")
    note = StudyPipeline(client).run(doc)
    assert note.title == "sample"
    assert len(note.sections) == 1
    assert len(client.calls) == 1
    # Document.content 进入 pipeline（user 消息包含正文片段）
    assert "普通段落内容" in client.calls[0][1].content

    out = save(note, tmp_path / "out")
    assert out.exists()
    text = out.read_text(encoding="utf-8")
    assert "# sample" in text
    assert "## 生成章节" in text


def test_e2e_pdf(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    _make_pdf(path)

    doc = parse_file(path)
    assert doc.format == "pdf"
    assert "--- Page 1 ---" in doc.content
    assert "--- Page 2 ---" in doc.content
    assert "第一页内容" in doc.content
    assert "第二页内容" in doc.content

    client = FakeLLMClient("## 生成章节")
    note = StudyPipeline(client).run(doc)
    assert note.title == "sample"
    assert len(note.sections) == 1
    assert "第一页内容" in client.calls[0][1].content

    out = save(note, tmp_path / "out")
    assert out.exists()
    assert "## 生成章节" in out.read_text(encoding="utf-8")


def test_e2e_pptx(tmp_path: Path) -> None:
    path = tmp_path / "sample.pptx"
    _make_pptx(path)

    doc = parse_file(path)
    assert doc.format == "pptx"
    # Slide 边界 / 标题 / bullet / 表格
    assert "--- Slide 1 ---" in doc.content
    assert "--- Slide 2 ---" in doc.content
    assert "# 第一页标题" in doc.content
    assert "- 要点一" in doc.content
    assert "| A | B |" in doc.content
    lines = [line.strip() for line in doc.content.split("\n")]
    assert "A" not in lines

    client = FakeLLMClient("## 生成章节")
    note = StudyPipeline(client).run(doc)
    assert note.title == "sample"
    assert len(note.sections) == 1
    assert "要点一" in client.calls[0][1].content

    out = save(note, tmp_path / "out")
    assert out.exists()
    assert "## 生成章节" in out.read_text(encoding="utf-8")


# ---------------------------------------------------------------------------
# CLI 端到端：main() 成功（exit 0）与失败（exit 1）
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "ext,maker",
    [(".docx", _make_docx), (".pdf", _make_pdf), (".pptx", _make_pptx)],
)
def test_cli_success_exit_zero(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, ext: str, maker: Any
) -> None:
    path = tmp_path / f"sample{ext}"
    maker(path)
    output_dir = tmp_path / "out"
    _patch_cli(monkeypatch)

    code = main([str(path), "-o", str(output_dir)])

    assert code == 0
    output = output_dir / "sample.md"
    assert output.exists()
    text = output.read_text(encoding="utf-8")
    assert "# sample" in text
    assert "## 生成章节" in text


def test_cli_unsupported_extension_exit_one(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "data.xyz"
    path.write_text("data", encoding="utf-8")
    _patch_cli(monkeypatch)

    code = main([str(path), "-o", str(tmp_path / "out")])

    assert code == 1


def test_cli_corrupted_file_exit_one(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "corrupted.docx"
    path.write_text("not a real docx", encoding="utf-8")
    _patch_cli(monkeypatch)

    code = main([str(path), "-o", str(tmp_path / "out")])

    assert code == 1


def test_cli_missing_file_exit_one(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "missing.docx"
    _patch_cli(monkeypatch)

    code = main([str(path), "-o", str(tmp_path / "out")])

    assert code == 1
