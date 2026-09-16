"""ParserRegistry + parse_file() 集成测试（真实文件路由，非 mock）。"""

from pathlib import Path

import pymupdf
import pytest
from docx import Document as _DocxDocument
from pptx import Presentation

from ai_study_assistant.parsers import (
    UnknownFormatError,
    default_registry,
    parse_file,
)


def _make_docx(path: Path) -> None:
    doc = _DocxDocument()
    doc.add_paragraph("docx integration content")
    doc.save(str(path))


def _make_pdf(path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox(pymupdf.Rect(72, 72, 500, 200), "pdf integration content")
    doc.save(str(path))
    doc.close()


def _make_pptx(path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "pptx integration content"
    prs.save(str(path))


def test_parse_file_docx(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    _make_docx(path)
    doc = parse_file(path)
    assert doc.format == "docx"
    assert "docx integration content" in doc.content


def test_parse_file_pdf(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    _make_pdf(path)
    doc = parse_file(path)
    assert doc.format == "pdf"
    assert "pdf integration content" in doc.content


def test_parse_file_pptx(tmp_path: Path) -> None:
    path = tmp_path / "sample.pptx"
    _make_pptx(path)
    doc = parse_file(path)
    assert doc.format == "pptx"
    assert "# pptx integration content" in doc.content


def test_parse_file_case_insensitive(tmp_path: Path) -> None:
    docx = tmp_path / "a.DOCX"
    _make_docx(docx)
    assert parse_file(docx).format == "docx"

    pdf = tmp_path / "b.PDF"
    _make_pdf(pdf)
    assert parse_file(pdf).format == "pdf"

    pptx = tmp_path / "c.PPTX"
    _make_pptx(pptx)
    assert parse_file(pptx).format == "pptx"


def test_parse_file_unknown_extension(tmp_path: Path) -> None:
    path = tmp_path / "data.xyz"
    path.write_text("data", encoding="utf-8")
    with pytest.raises(UnknownFormatError):
        parse_file(path)


def test_supported_extensions() -> None:
    assert default_registry.supported_extensions() == {
        ".md",
        ".markdown",
        ".txt",
        ".docx",
        ".pdf",
        ".pptx",
    }


def test_markdown_text_routing_unchanged(tmp_path: Path) -> None:
    md = tmp_path / "note.md"
    md.write_text("# Hello", encoding="utf-8")
    assert parse_file(md).format == "markdown"

    markdown = tmp_path / "note.markdown"
    markdown.write_text("content", encoding="utf-8")
    assert parse_file(markdown).format == "markdown"

    txt = tmp_path / "note.txt"
    txt.write_text("plain text", encoding="utf-8")
    assert parse_file(txt).format == "text"
